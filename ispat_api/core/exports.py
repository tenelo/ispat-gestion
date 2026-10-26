"""
Exports Excel (.xlsx) et CSV communs à tous les modules, aux couleurs ISPAT.

    return reponse_export("xlsx", "Liste de classe", "L1 Agronomie A — 2026-2027",
                          ["N°", "Matricule", "Nom"], lignes, "liste_L1_AGR_A")
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VERT_ISPAT = "1B6B3A"
VERT_CLAIR = "EAF3EC"
ORANGE_ISPAT = "C8792A"
FORMATS_EXPORT = ("xlsx", "csv")


def nom_etablissement():
    from referentiels.models import Parametre      # import local : évite une dépendance circulaire au chargement

    parametre = Parametre.objects.filter(cle="nom_etablissement").first()
    return parametre.valeur if parametre else "ISPAT"


def nom_fichier_sur(nom):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", nom).strip("_")[:80] or "export"


def _valeur_cellule(valeur):
    if isinstance(valeur, Decimal):
        return float(valeur)
    return valeur


def _texte(valeur):
    if valeur is None:
        return ""
    if isinstance(valeur, datetime):
        return timezone.localtime(valeur).strftime("%d/%m/%Y %H:%M") if timezone.is_aware(valeur) else valeur.strftime("%d/%m/%Y %H:%M")
    if isinstance(valeur, date):
        return valeur.strftime("%d/%m/%Y")
    return str(valeur)


def reponse_export(format_export, titre, sous_titre, entetes, lignes, nom_fichier):
    nom = nom_fichier_sur(nom_fichier)
    if format_export == "csv":
        tampon = io.StringIO()
        ecrivain = csv.writer(tampon, delimiter=";")          # séparateur attendu par Excel en français
        ecrivain.writerow(entetes)
        for ligne in lignes:
            ecrivain.writerow([_texte(v) for v in ligne])
        reponse = HttpResponse("﻿" + tampon.getvalue(), content_type="text/csv; charset=utf-8")
        reponse["Content-Disposition"] = f'attachment; filename="{nom}.csv"'
        return reponse

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = nom_fichier_sur(titre)[:31] or "Export"
    nb_colonnes = max(len(entetes), 1)
    derniere = get_column_letter(nb_colonnes)

    feuille.merge_cells(f"A1:{derniere}1")
    feuille["A1"] = nom_etablissement()
    feuille["A1"].font = Font(bold=True, size=12, color=VERT_ISPAT)
    feuille.merge_cells(f"A2:{derniere}2")
    feuille["A2"] = titre
    feuille["A2"].font = Font(bold=True, size=14)
    feuille.merge_cells(f"A3:{derniere}3")
    feuille["A3"] = f"{sous_titre} — édité le {timezone.localtime().strftime('%d/%m/%Y à %H:%M')}" if sous_titre \
        else f"Édité le {timezone.localtime().strftime('%d/%m/%Y à %H:%M')}"
    feuille["A3"].font = Font(italic=True, size=9, color=ORANGE_ISPAT)

    trait = Side(style="thin", color="B7C9BC")
    bordure = Border(left=trait, right=trait, top=trait, bottom=trait)
    ligne_entete = 5
    for col, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=col, value=entete)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor=VERT_ISPAT)
        cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cellule.border = bordure

    largeurs = [len(str(e)) for e in entetes]
    for i, ligne in enumerate(lignes):
        for col, valeur in enumerate(ligne, start=1):
            cellule = feuille.cell(row=ligne_entete + 1 + i, column=col, value=_valeur_cellule(valeur))
            cellule.border = bordure
            if isinstance(valeur, date):
                cellule.number_format = "DD/MM/YYYY"
            if i % 2:
                cellule.fill = PatternFill("solid", fgColor=VERT_CLAIR)
            largeurs[col - 1] = max(largeurs[col - 1], len(_texte(valeur)))

    for col, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(col)].width = min(max(largeur + 2, 6), 45)
    feuille.freeze_panes = feuille.cell(row=ligne_entete + 1, column=1)
    if lignes:
        feuille.auto_filter.ref = f"A{ligne_entete}:{derniere}{ligne_entete + len(lignes)}"
    feuille.page_setup.orientation = "landscape" if nb_colonnes > 7 else "portrait"
    feuille.page_setup.fitToWidth = 1
    feuille.sheet_properties.pageSetUpPr.fitToPage = True
    feuille.print_title_rows = f"{ligne_entete}:{ligne_entete}"

    tampon = io.BytesIO()
    classeur.save(tampon)
    reponse = HttpResponse(tampon.getvalue(),
                           content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    reponse["Content-Disposition"] = f'attachment; filename="{nom}.xlsx"'
    return reponse
